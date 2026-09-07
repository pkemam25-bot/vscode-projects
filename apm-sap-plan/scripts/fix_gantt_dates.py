"""
Fix Gantt chart in apm-wbs-template-v4.xlsx

Root cause: Columns E (Planned Start Date) and F (Planned End Date) contain
inline strings (text like "2026-08-21") instead of proper Excel date serials.
The Gantt formulas =IF(AND($E5<=DATE(...), $F5>=DATE(...)), "■", "")
always evaluate to FALSE because they compare text to numbers.

This script converts all date cells in E5:F36 to proper datetime objects
so the formulas evaluate correctly and the Gantt bars appear.
"""

import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl
from datetime import datetime, date

FILE_IN  = r"../wbs-templates/apm-wbs-template-v4.xlsx"
FILE_OUT = r"../wbs-templates/apm-wbs-template-v4-fixed.xlsx"

def parse_date(value):
    """Parse a date string into a date object."""
    if isinstance(value, (datetime, date)):
        return value
    if not isinstance(value, str):
        return None
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def main():
    wb = openpyxl.load_workbook(FILE_IN)
    ws = wb["APM WBS Schedule"]

    converted = 0
    errors = []

    for row in range(5, ws.max_row + 1):
        for col_letter in ["E", "F"]:
            cell = ws[f"{col_letter}{row}"]
            raw = cell.value
            if raw is None or raw == "":
                continue

            d = parse_date(raw)
            if d is None:
                errors.append(f"{col_letter}{row}: cannot parse {repr(raw)}")
                continue

            cell.value = d
            converted += 1

    wb.save(FILE_OUT)

    print(f"Converted {converted} date cells to proper Excel dates")
    if errors:
        print(f"\nErrors ({len(errors)}):")
        for e in errors:
            print(f"  {e}")
    else:
        print("No errors")

    print(f"\nSaved to: {FILE_OUT}")

    # Quick verification: re-open and check types
    wb2 = openpyxl.load_workbook(FILE_OUT)
    ws2 = wb2["APM WBS Schedule"]
    print("\nVerification (sample cells):")
    for ref in ["E5", "F5", "E10", "F10", "E25", "F25", "E33", "F33"]:
        cell = ws2[ref]
        print(f"  {ref}: value={repr(cell.value)}, type={type(cell.value).__name__}")

    # Verify a formula would work
    e5 = ws2["E5"].value
    f5 = ws2["F5"].value
    print(f"\nFormula check for row 5:")
    print(f"  E5 = {e5} (type: {type(e5).__name__})")
    print(f"  F5 = {f5} (type: {type(f5).__name__})")
    if isinstance(e5, date) and isinstance(f5, date):
        # Simulate: M5 = IF(AND($E5<=DATE(2026,8,31), $F5>=DATE(2026,8,1)), "■", "")
        aug_start = datetime(2026, 8, 1)
        aug_end = datetime(2026, 8, 31)
        result = "■" if (e5 <= aug_end and f5 >= aug_start) else ""
        print(f"  M5 (Aug '26): E5 <= {aug_end} AND F5 >= {aug_start} => {repr(result)}")
        print("  ✅ Formula logic is now correct!")
    else:
        print("  ❌ Dates are not proper date objects")


if __name__ == "__main__":
    main()
