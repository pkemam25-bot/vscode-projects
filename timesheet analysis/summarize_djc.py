import openpyxl
from datetime import datetime
from collections import defaultdict

# Load the workbook
wb = openpyxl.load_workbook('TIME_report_25_Aug.xlsx', data_only=True)
ws = wb[wb.sheetnames[0]]

# Extract all Digital Job Card entries for Apr-May
digital_job_card_data = []

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
    package = row[2].value  # Column C = Package
    date_val = row[4].value  # Column E = Date
    hours = row[5].value  # Column F = Hours
    activity = row[3].value  # Column D = Activity description
    
    if package and 'Digital Job Card' in str(package):
        if isinstance(date_val, datetime):
            month = date_val.month
            day = date_val.day
            year = date_val.year
            # Filter for Apr (4) and May (5) only
            if month in [4, 5]:
                digital_job_card_data.append({
                    'year': year,
                    'month': month,
                    'day': day,
                    'date': date_val,
                    'hours': hours if hours else 0,
                    'activity': activity,
                    'date_str': date_val.strftime('%Y-%m-%d')
                })

# Sort by date
digital_job_card_data.sort(key=lambda x: x['date'])

# Print summary
print(f"Total records found: {len(digital_job_card_data)}")
print()

# Group by month
monthly_totals = defaultdict(float)
daily_data = defaultdict(lambda: defaultdict(float))

for entry in digital_job_card_data:
    month_name = entry['date'].strftime('%B')
    month_num = entry['month']
    date_str = entry['date'].strftime('%Y-%m-%d')
    
    daily_data[month_num][date_str] += entry['hours']
    monthly_totals[month_num] += entry['hours']

# Print table
print("=" * 50)
print("Digital Job Card - Time Summary (Apr-May)")
print("=" * 50)
print()
print(f"{'Month':<12} {'Date':<15} {'Time (hrs)':<12}")
print("-" * 40)

for month_num in [4, 5]:
    if month_num in daily_data:
        month_name = datetime(2026, month_num, 1).strftime('%B')
        for date_str in sorted(daily_data[month_num].keys()):
            hours = daily_data[month_num][date_str]
            print(f"{month_name:<12} {date_str:<15} {hours:<12.1f}")
        # Monthly subtotal
        print(f"{month_name} Total: {monthly_totals[month_num]:.1f} hrs")
        print()

# Grand total
grand_total = sum(monthly_totals.values())
print(f"{'Grand Total:':<12} {grand_total:.1f} hrs")
